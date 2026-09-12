import openpyxl
import pandas as pd
from datetime import datetime
import io

def process_chalimeda_feeds(file_bytes):
    # Load the workbook directly from memory instead of a folder path
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)

    IGNORED_EXACT = {"MIS", "O.STOCK", "SHEET1", "SHEET2", "SHEET3", "SHEET4", "SHEET5"}
    CLOSING_KEYWORDS = [
        "closing", "c.stock", "c/stock", "cl stock", "cl.stock", "bal stock",
        "c stock", "cl. stock", "cb", "c/b", "c.b", "balance", "bal", "c.bal",
        "cl.bal", "closing bal", "closing stock", "closing qty", "stock"
    ]
    DATE_KEYWORDS = ["date", "dt", "day", "particulars", "sl.no", "sl no", "s.no"]

    def get_cell_val_with_merge_support(sheet, row, col):
        val = sheet.cell(row=row, column=col).value
        if val is not None:
            return str(val).strip()

        for rng in sheet.merged_cells.ranges:
            if row >= rng.min_row and row <= rng.max_row and col >= rng.min_col and col <= rng.max_col:
                top_left_val = sheet.cell(row=rng.min_row, column=rng.min_col).value
                if top_left_val is not None:
                    return str(top_left_val).strip()
        return ""

    def find_sheet_columns(sheet):
        date_col = None
        closing_col = None
        header_row = None

        for r in range(1, 25):
            for c in range(1, 20):
                val1 = get_cell_val_with_merge_support(sheet, r, c).lower()
                val0 = get_cell_val_with_merge_support(sheet, max(1, r-1), c).lower()
                combined = f"{val0} {val1}".strip()

                if not date_col and any(k in combined for k in DATE_KEYWORDS):
                    date_col = c
                    header_row = r

                if not closing_col and any(k in combined for k in CLOSING_KEYWORDS):
                    if "opening" in combined and "closing" not in combined:
                        continue
                    closing_col = c
                    header_row = header_row or r

            if date_col and closing_col:
                break

        if not date_col:
            date_col = 1
        return header_row, date_col, closing_col

    results = []
    for sheet_name in wb.sheetnames:
        clean_name = sheet_name.strip()
        if clean_name.upper() in IGNORED_EXACT:
            continue

        sheet = wb[sheet_name]
        header_row, date_col, closing_col = find_sheet_columns(sheet)

        if not closing_col:
            continue

        latest_date = None
        latest_stock = None
        start_row = (header_row or 1) + 1

        for r in range(start_row, sheet.max_row + 1):
            date_val = sheet.cell(row=r, column=date_col).value
            stock_val = sheet.cell(row=r, column=closing_col).value

            if stock_val is None or str(stock_val).strip() == "":
                merged_val = get_cell_val_with_merge_support(sheet, r, closing_col)
                if merged_val:
                    stock_val = merged_val

            if stock_val is not None and str(stock_val).strip() != "":
                if isinstance(date_val, datetime):
                    fmt_date = date_val.strftime("%d-%b-%Y")
                elif date_val is not None and str(date_val).strip():
                    fmt_date = str(date_val).strip()
                else:
                    fmt_date = "N/A"

                latest_date = fmt_date

                try:
                    num_val = float(stock_val)
                    num_val = round(num_val, 3)
                    if abs(num_val) < 0.0001:
                        num_val = 0.0
                    latest_stock = f"{num_val:.3f}"
                except (ValueError, TypeError):
                    latest_stock = str(stock_val).strip()

        if latest_stock is not None:
            results.append({
                "Raw Material": clean_name,
                "Latest Date": latest_date,
                "Closing Stock": latest_stock
            })

    # Format into WhatsApp Message
    whatsapp_msg = f"📊 *DAILY RAW MATERIAL CLOSING STOCKS*\n📅 Date: {results[0]['Latest Date'] if results else 'N/A'}\n" + "—"*30 + "\n"
    for item in results:
        whatsapp_msg += f"• *{item['Raw Material']}*: {item['Closing Stock']}\n"

    # Return the data as a JSON dictionary to n8n
    return {
        "success": True,
        "total_processed": len(results),
        "whatsapp_message": whatsapp_msg,
        "data": results
    }
